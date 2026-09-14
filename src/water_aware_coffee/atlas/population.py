"""Official resident-population figures for the localities in the atlas.

Two raw sources (data/raw/population/):
- ibge/sidra_t4709_v93_n6_p2022.json: IBGE Censo Demográfico 2022, SIDRA table 4709 "População
  residente" (variable 93), territorial level N6 (município), period 2022, as returned by the SIDRA
  API. A JSON list whose first element is the header row (column labels); each data row carries the
  7-digit municipality code in ``D1C`` and the count in ``V``. SISAGUA keys municipalities by the
  first 6 digits of that code (the 7th is a check digit), so ``locality_code`` is ``D1C[:6]``.
- ons/sapelsoasyoa*.xlsx: ONS "Lower layer Super Output Area population estimates", England and
  Wales, one worksheet per mid-year ("Mid-2024 LSOA 2021"). The header row is the one whose third
  cell reads "LSOA 2021 Code"; the "Total" column is the all-ages, both-sexes estimate. The latest
  mid-year sheet in the workbook is used and its year is recorded in ``population_year``.

``load_population`` reads whichever of the two is present; ``attach_population`` fills the atlas
``population`` column where it is null by matching ``(country_iso2, locality_key)`` and leaves
values that came with a source (US, from EPA SYR4) untouched.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
import pandas as pd

IBGE_SUBDIR = "ibge"
ONS_SUBDIR = "ons"
IBGE_GLOB = "sidra_t4709*.json"
ONS_GLOB = "sapelsoasyoa*.xlsx"
IBGE_SOURCE = "ibge-censo-2022-sidra-t4709"
ONS_SOURCE = "ons-lsoa-mid-year-estimates"
IBGE_YEAR = 2022

POPULATION_COLUMNS = [
    "country_iso2",
    "locality_code",
    "population",
    "population_source",
    "population_year",
]

_ONS_SHEET = re.compile(r"^Mid-(\d{4}) LSOA 2021$")
_ONS_CODE_HEADER = "LSOA 2021 Code"
_ONS_TOTAL_HEADER = "Total"
_LSOA_CODE = re.compile(r"^[EW]01\d{6}$")


def _empty() -> pd.DataFrame:
    return pd.DataFrame(columns=POPULATION_COLUMNS).astype(
        {
            "country_iso2": "string",
            "locality_code": "string",
            "population": "Int64",
            "population_source": "string",
            "population_year": "Int64",
        }
    )


def load_ibge(path: Path) -> pd.DataFrame:
    """Municipality population from one SIDRA API JSON response (table 4709, variable 93)."""
    with path.open(encoding="utf-8") as fh:
        records = json.load(fh)
    if not records or records[0].get("D1C") != "Município (Código)":
        raise ValueError(f"{path}: not a SIDRA table-4709 municipality response")
    df = pd.DataFrame(records[1:])
    if not (df["D2C"] == "93").all():
        raise ValueError(f"{path}: expected only variable 93 (População residente)")
    code7 = df["D1C"].astype("string").str.strip()
    if not code7.str.fullmatch(r"\d{7}").all():
        raise ValueError(f"{path}: municipality codes are not all 7 digits")
    out = pd.DataFrame(
        {
            "country_iso2": "BR",
            "locality_code": code7.str[:6],
            "population": pd.to_numeric(df["V"], errors="coerce"),
            "population_source": IBGE_SOURCE,
            "population_year": pd.to_numeric(df["D3C"], errors="coerce"),
        }
    )
    out = out[out["population"].notna()]
    return _finish(out)


def _latest_ons_sheet(names: list[str]) -> tuple[str, int]:
    matches = [(int(m.group(1)), n) for n in names if (m := _ONS_SHEET.match(n))]
    if not matches:
        raise ValueError(f"no 'Mid-YYYY LSOA 2021' worksheet among {names}")
    year, name = max(matches)
    return name, year


def load_ons(path: Path) -> pd.DataFrame:
    """LSOA (2021 boundaries) all-ages population from the latest mid-year sheet in the workbook."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet, year = _latest_ons_sheet(wb.sheetnames)
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        header: tuple[object, ...] | None = None
        for row in rows:
            if _ONS_CODE_HEADER in row:
                header = row
                break
        if header is None:
            raise ValueError(f"{path}[{sheet}]: header row with '{_ONS_CODE_HEADER}' not found")
        code_i = header.index(_ONS_CODE_HEADER)
        total_i = header.index(_ONS_TOTAL_HEADER)
        data = [
            (row[code_i], row[total_i])
            for row in rows
            if len(row) > max(code_i, total_i) and row[code_i] is not None
        ]
    finally:
        wb.close()
    df = pd.DataFrame(data, columns=["locality_code", "population"])
    df["locality_code"] = df["locality_code"].astype("string").str.strip()
    df = df[df["locality_code"].str.fullmatch(_LSOA_CODE.pattern)]
    out = pd.DataFrame(
        {
            "country_iso2": "GB",
            "locality_code": df["locality_code"],
            "population": pd.to_numeric(df["population"], errors="coerce"),
            "population_source": ONS_SOURCE,
            "population_year": year,
        }
    )
    out = out[out["population"].notna()]
    return _finish(out)


def _finish(df: pd.DataFrame) -> pd.DataFrame:
    out = df.loc[:, POPULATION_COLUMNS].copy()
    out["country_iso2"] = out["country_iso2"].astype("string")
    out["locality_code"] = out["locality_code"].astype("string")
    out["population"] = out["population"].round().astype("Int64")
    out["population_source"] = out["population_source"].astype("string")
    out["population_year"] = out["population_year"].astype("Int64")
    dupes = out.duplicated(["country_iso2", "locality_code"], keep=False)
    if dupes.any():
        raise ValueError(
            f"duplicate locality codes: {out.loc[dupes, 'locality_code'].tolist()[:5]}"
        )
    return out.reset_index(drop=True)


def load_population(raw_dir: Path) -> pd.DataFrame:
    """One row per (country_iso2, locality_code) from every population file under ``raw_dir``.

    ``raw_dir`` is data/raw/population; the IBGE and ONS files are looked up in its ``ibge`` and
    ``ons`` subdirectories. Missing subdirectories are skipped; raises if nothing is found.
    """
    parts: list[pd.DataFrame] = []
    for path in sorted((raw_dir / IBGE_SUBDIR).glob(IBGE_GLOB)):
        parts.append(load_ibge(path))
    for path in sorted((raw_dir / ONS_SUBDIR).glob(ONS_GLOB)):
        parts.append(load_ons(path))
    if not parts:
        raise FileNotFoundError(f"no population files under {raw_dir}")
    return _finish(pd.concat(parts, ignore_index=True))


def attach_population(wide: pd.DataFrame, pop: pd.DataFrame) -> pd.DataFrame:
    """Fill null ``population`` in the atlas wide table from ``pop``; keep existing values.

    Matches on ``(country_iso2, locality_key) == (country_iso2, locality_code)``. Adds
    ``population_source``: the atlas ``source_id`` for rows whose population came with the
    measurements, the population dataset id for rows filled here, null where still unknown.
    ``population_year`` is added for filled rows as well.
    """
    out = wide.copy()
    had = out["population"].notna()
    if "population_source" not in out.columns:
        out["population_source"] = pd.Series(pd.NA, index=out.index, dtype="string")
    if "population_year" not in out.columns:
        out["population_year"] = pd.Series(pd.NA, index=out.index, dtype="Int64")
    out.loc[had & out["population_source"].isna(), "population_source"] = out.loc[
        had & out["population_source"].isna(), "source_id"
    ]
    lookup = pop.set_index(["country_iso2", "locality_code"])
    keys = pd.MultiIndex.from_arrays(
        [out["country_iso2"].astype("string"), out["locality_key"].astype("string")]
    )
    matched = lookup.reindex(keys)
    fill = ~had & matched["population"].notna().to_numpy()
    out.loc[fill, "population"] = matched["population"].to_numpy()[fill]
    out.loc[fill, "population_source"] = matched["population_source"].to_numpy()[fill]
    out.loc[fill, "population_year"] = matched["population_year"].to_numpy()[fill]
    out["population"] = out["population"].astype("Int64")
    return out
